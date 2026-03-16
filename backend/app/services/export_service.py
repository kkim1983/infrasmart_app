import io
from datetime import datetime
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
SUB_HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
SUB_HEADER_FONT = Font(bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_header(ws, row, col, value, width=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER_ALIGN
    cell.border = THIN_BORDER
    return cell


def _apply_cell(ws, row, col, value, bold=False, center=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, size=10)
    cell.alignment = CENTER_ALIGN if center else Alignment(vertical="center", wrap_text=True)
    cell.border = THIN_BORDER
    return cell


class ExportService:
    """엑셀 보고서 생성 서비스"""

    def generate_inspection_report(
        self,
        inspection: dict,
        facility: dict,
        inspector: dict,
        damage_records: list[dict],
        photos: list[dict],
        gps_track: list[dict],
    ) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)  # 기본 시트 제거

        self._sheet_overview(wb, inspection, facility, inspector)
        self._sheet_damage_table(wb, damage_records, facility)
        self._sheet_photo_layout(wb, photos)
        self._sheet_auth_history(wb, inspection, inspector, gps_track)
        self._sheet_gps_track(wb, gps_track)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # ── Sheet 1: 현장개요 ─────────────────────────────────────
    def _sheet_overview(self, wb, inspection, facility, inspector):
        ws = wb.create_sheet("1. 현장개요")
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 30

        # 제목
        ws.merge_cells("A1:D1")
        title = ws["A1"]
        title.value = f"안전점검 결과보고서 - {facility.get('name', '')}"
        title.font = Font(bold=True, size=14)
        title.alignment = CENTER_ALIGN
        title.fill = PatternFill("solid", fgColor="1F4E79")
        title.font = Font(color="FFFFFF", bold=True, size=14)

        rows = [
            ("시설물 코드", facility.get("code", ""), "시설물 명칭", facility.get("name", "")),
            ("시설물 유형", facility.get("facility_type", ""), "소재지", facility.get("location_name", "")),
            ("점검 유형", inspection.get("inspection_type", ""), "점검 상태", inspection.get("status", "")),
            ("점검 시작일", str(inspection.get("started_at", "")), "점검 종료일", str(inspection.get("ended_at", ""))),
            ("점검자 성명", inspector.get("name", ""), "자격증 번호", inspector.get("license_no", "")),
            ("자격 종류", inspector.get("license_type", ""), "연락처", inspector.get("phone", "")),
        ]

        meta = facility.get("facility_meta") or {}
        rows += [
            ("교량 연장(m)", meta.get("bridge_length", ""), "교량 폭원(m)", meta.get("width", "")),
            ("스팬 수", meta.get("span_count", ""), "준공년도", meta.get("built_year", "")),
        ]

        for i, (k1, v1, k2, v2) in enumerate(rows, start=2):
            _apply_cell(ws, i, 1, k1, bold=True, center=True).fill = SUB_HEADER_FILL
            _apply_cell(ws, i, 2, v1)
            _apply_cell(ws, i, 3, k2, bold=True, center=True).fill = SUB_HEADER_FILL
            _apply_cell(ws, i, 4, v2)

    # ── Sheet 2: 손상현황표 ───────────────────────────────────
    def _sheet_damage_table(self, wb, damage_records, facility):
        ws = wb.create_sheet("2. 손상현황표")
        headers = [
            "번호", "손상코드", "손상명칭", "위치", "규모",
            "상태등급", "신규여부", "확대여부", "보수여부", "보수일자", "비고"
        ]
        col_widths = [6, 10, 14, 20, 20, 8, 8, 8, 8, 12, 20]

        for col, (h, w) in enumerate(zip(headers, col_widths), 1):
            ws.column_dimensions[get_column_letter(col)].width = w
            _apply_header(ws, 1, col, h)

        for row, rec in enumerate(damage_records, 2):
            loc = rec.get("location_data") or {}
            dims = rec.get("dimensions") or {}
            location_str = (
                f"부재:{loc.get('member_no','')} 스팬:{loc.get('span_no','')}"
                if loc else ""
            )
            dimension_str = " ".join(f"{k}={v}" for k, v in dims.items())

            values = [
                row - 1,
                rec.get("damage_code", ""),
                rec.get("damage_name", ""),
                location_str,
                dimension_str,
                rec.get("severity_grade", ""),
                "신규" if rec.get("is_new") else "기존",
                "확대" if rec.get("is_expanded") else "-",
                "완료" if rec.get("is_repaired") else "-",
                str(rec.get("repair_date", "") or ""),
                rec.get("notes", ""),
            ]
            for col, val in enumerate(values, 1):
                _apply_cell(ws, row, col, val, center=(col <= 6))

    # ── Sheet 3: 사진대지 ─────────────────────────────────────
    def _sheet_photo_layout(self, wb, photos):
        ws = wb.create_sheet("3. 사진대지")
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 40

        _apply_header(ws, 1, 1, "사진번호")
        _apply_header(ws, 1, 2, "손상 내용 및 위치")
        _apply_header(ws, 1, 3, "촬영일시")
        _apply_header(ws, 1, 4, "GPS 좌표")

        for row, photo in enumerate(photos, 2):
            gps = photo.get("gps") or {}
            gps_str = f"{gps.get('lat','')}, {gps.get('lng','')}" if gps else ""
            _apply_cell(ws, row, 1, photo.get("photo_number", ""), center=True)
            _apply_cell(ws, row, 2, photo.get("damage_description", ""))
            _apply_cell(ws, row, 3, str(photo.get("taken_at", "") or ""), center=True)
            _apply_cell(ws, row, 4, gps_str, center=True)
            ws.row_dimensions[row].height = 80  # 사진 삽입 공간

    # ── Sheet 4: 점검자 인증이력 ──────────────────────────────
    def _sheet_auth_history(self, wb, inspection, inspector, gps_track):
        ws = wb.create_sheet("4. 인증이력")
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 40

        _apply_header(ws, 1, 1, "항목")
        _apply_header(ws, 1, 2, "내용")

        rows = [
            ("점검자 성명", inspector.get("name", "")),
            ("자격증 번호", inspector.get("license_no", "")),
            ("자격 종류", inspector.get("license_type", "")),
            ("점검 시작", str(inspection.get("started_at", ""))),
            ("점검 종료", str(inspection.get("ended_at", ""))),
            ("GPS 기록 수", str(len(gps_track))),
            ("기기 정보", str(inspection.get("device_info", {}))),
        ]

        for i, (k, v) in enumerate(rows, 2):
            _apply_cell(ws, i, 1, k, bold=True, center=True).fill = SUB_HEADER_FILL
            _apply_cell(ws, i, 2, v)

    # ── Sheet 5: GPS 이력 ─────────────────────────────────────
    def _sheet_gps_track(self, wb, gps_track):
        ws = wb.create_sheet("5. GPS이력")
        headers = ["순번", "위도", "경도", "정확도(m)", "시간"]
        col_widths = [6, 16, 16, 12, 25]

        for col, (h, w) in enumerate(zip(headers, col_widths), 1):
            ws.column_dimensions[get_column_letter(col)].width = w
            _apply_header(ws, 1, col, h)

        for row, point in enumerate(gps_track, 2):
            _apply_cell(ws, row, 1, row - 1, center=True)
            _apply_cell(ws, row, 2, point.get("lat", ""), center=True)
            _apply_cell(ws, row, 3, point.get("lng", ""), center=True)
            _apply_cell(ws, row, 4, point.get("accuracy", ""), center=True)
            _apply_cell(ws, row, 5, str(point.get("time", "") or ""), center=True)


export_service = ExportService()
